"""
Unit tests for formatter.py and excel_writer.py modules.
"""

import unittest
import os
import sys

# Ensure parent directory is in sys.path
BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if BASE_DIR not in sys.path:
    sys.path.insert(0, BASE_DIR)

import openpyxl
import config
from formatter import convert_otis_report


class TestFormatter(unittest.TestCase):

    def test_end_to_end_conversion(self):
        sample_path = os.path.join(config.SAMPLE_DIR, "sample_test_sequence.xlsx")
        output_path = os.path.join(config.OUTPUT_DIR, "test_output.xlsx")

        if os.path.exists(output_path):
            os.remove(output_path)

        saved_path, blocks, summary = convert_otis_report(sample_path, output_path)

        self.assertTrue(os.path.exists(saved_path))
        self.assertEqual(summary["total_reports"], len(blocks))
        # Check row count: (5 blocks * 5 rows) + 1 header row = 26 rows
        self.assertEqual(summary["total_rows_written"], (len(blocks) * 5) + 1)

        # Inspect generated openpyxl workbook
        wb = openpyxl.load_workbook(saved_path)
        self.assertIn("OTIS Inspection Report", wb.sheetnames)
        ws = wb["OTIS Inspection Report"]

        # Check Header
        headers = [ws.cell(1, c).value for c in range(1, 8)]
        self.assertEqual(headers, config.OUTPUT_COLUMNS)

        # Check block 1 merged cells (A2:A6, B2:B6)
        merged_ranges = [str(m) for m in ws.merged_cells.ranges]
        self.assertIn("A2:A6", merged_ranges)
        self.assertIn("B2:B6", merged_ranges)

        # Cleanup
        wb.close()
        if os.path.exists(output_path):
            os.remove(output_path)


if __name__ == "__main__":
    unittest.main()
