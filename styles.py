"""
Styles module for OTIS Excel Report Formatter.
Provides openpyxl Font, Alignment, Border, and Fill definitions
matching OTIS/Kaynes report layout standards.
"""

from openpyxl.styles import Font, Alignment, Border, Side, PatternFill
from openpyxl.worksheet.worksheet import Worksheet
import config

# Colors
HEADER_FILL_COLOR = config.HEADER_FILL_COLOR
BORDER_COLOR_HEX = config.BORDER_COLOR

# Fonts
HEADER_FONT = Font(
    name=config.FONT_FAMILY,
    size=config.HEADER_FONT_SIZE,
    bold=True,
    color="000000"
)

BODY_FONT = Font(
    name=config.FONT_FAMILY,
    size=config.BODY_FONT_SIZE,
    bold=False,
    color="000000"
)

BODY_BOLD_FONT = Font(
    name=config.FONT_FAMILY,
    size=config.BODY_FONT_SIZE,
    bold=True,
    color="000000"
)

BODY_ITALIC_FONT = Font(
    name=config.FONT_FAMILY,
    size=config.BODY_FONT_SIZE,
    italic=True,
    color="000000"
)

# Fills
HEADER_FILL = PatternFill(
    start_color=HEADER_FILL_COLOR,
    end_color=HEADER_FILL_COLOR,
    fill_type="solid"
)

# Borders
THIN_SIDE = Side(border_style="thin", color=BORDER_COLOR_HEX)
MEDIUM_SIDE = Side(border_style="medium", color=BORDER_COLOR_HEX)

THIN_BORDER = Border(
    left=THIN_SIDE,
    right=THIN_SIDE,
    top=THIN_SIDE,
    bottom=THIN_SIDE
)

HEADER_BORDER = Border(
    left=THIN_SIDE,
    right=THIN_SIDE,
    top=MEDIUM_SIDE,
    bottom=MEDIUM_SIDE
)

# Alignments
HEADER_ALIGNMENT = Alignment(
    horizontal="center",
    vertical="center",
    wrap_text=True
)

BODY_LEFT_ALIGNMENT = Alignment(
    horizontal="left",
    vertical="center",
    wrap_text=True
)

BODY_CENTER_ALIGNMENT = Alignment(
    horizontal="center",
    vertical="center",
    wrap_text=True
)

BODY_RIGHT_ALIGNMENT = Alignment(
    horizontal="right",
    vertical="center",
    wrap_text=True
)


def style_cell(cell, font=None, fill=None, border=None, alignment=None):
    """Applies specified openpyxl styles to a single cell."""
    if font:
        cell.font = font
    if fill:
        cell.fill = fill
    if border:
        cell.border = border
    if alignment:
        cell.alignment = alignment


def apply_range_borders(ws: Worksheet, start_col: int, start_row: int, end_col: int, end_row: int, border: Border = THIN_BORDER) -> None:
    """
    Applies border to every cell within a rectangular range.
    Ensures merged cell ranges maintain full borders.
    """
    for r in range(start_row, end_row + 1):
        for c in range(start_col, end_col + 1):
            cell = ws.cell(row=r, column=c)
            cell.border = border
