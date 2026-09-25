"""
Excel Writer module for OTIS Excel Report Formatter.
Constructs the final OTIS Inspection Report Excel workbook using openpyxl,
applying 5-row structured layouts, cell merging, styling, and page setup.
"""

from typing import List, Union, Dict, Any, Optional
import openpyxl
from openpyxl.worksheet.worksheet import Worksheet
import config
from logger import logger
from parser import ReportBlock
from styles import (
    HEADER_FONT, BODY_FONT, BODY_ITALIC_FONT,
    HEADER_FILL, THIN_BORDER, HEADER_BORDER,
    HEADER_ALIGNMENT, BODY_LEFT_ALIGNMENT, BODY_CENTER_ALIGNMENT,
    apply_range_borders
)
from utils import format_number_display


def build_otis_inspection_workbook(report_blocks: List[ReportBlock]) -> openpyxl.Workbook:
    """
    Builds the formatted OTIS Inspection Report openpyxl Workbook from parsed ReportBlock objects.
    Each report block generates exactly 5 formatted rows with merged Sl No and Test Name columns.
    """
    logger.info("Building formatted OTIS Inspection Report workbook...")
    wb = openpyxl.Workbook()
    ws: Worksheet = wb.active
    ws.title = "OTIS Inspection Report"

    # Write Header Row
    ws.append(config.OUTPUT_COLUMNS)
    ws.row_dimensions[1].height = config.HEADER_ROW_HEIGHT

    for col_idx, col_name in enumerate(config.OUTPUT_COLUMNS, 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.border = HEADER_BORDER
        cell.alignment = HEADER_ALIGNMENT

    current_row = 2

    for block in report_blocks:
        start_row = current_row

        # Prepare formatted values
        v_str = format_number_display(block.voltage, "V") if block.voltage else block.voltage_raw
        i_str = format_number_display(block.current, "A") if block.current else block.current_raw

        set_v_val = f"SET Voltage {v_str}".strip() if v_str else "SET Voltage"
        set_i_val = f"SET CURRENT {i_str}".strip() if i_str else "SET CURRENT"

        read_v_text = f"read voltage {block.dmm_point}".strip() if block.dmm_point else "read voltage"

        min_str = block.min_voltage if block.min_voltage is not None else "-"
        max_str = block.max_voltage if block.max_voltage is not None else "-"

        # Row 1: PSU - Battery Point
        r1 = [
            block.report_no,
            block.test_name,
            block.psu_device,
            block.battery_point if block.battery_point else "P_BAT",
            "-",
            "-",
            ""
        ]

        # Row 2: PSU - SET Voltage
        r2 = [
            "",
            "",
            block.psu_device,
            set_v_val,
            "-",
            "-",
            ""
        ]

        # Row 3: PSU - SET CURRENT
        r3 = [
            "",
            "",
            block.psu_device,
            set_i_val,
            "-",
            "-",
            ""
        ]

        # Row 4: Waveshare - Measurement Point
        r4 = [
            "",
            "",
            block.waveshare_device,
            block.measurement_point if block.measurement_point else "Measurement Point",
            "-",
            "-",
            ""
        ]

        # Row 5: DMM - read voltage & Min/Max limits
        r5 = [
            "",
            "",
            block.dmm_device,
            read_v_text,
            min_str,
            max_str,
            ""
        ]

        block_rows = [r1, r2, r3, r4, r5]

        for r_offset, r_data in enumerate(block_rows):
            r_num = start_row + r_offset
            ws.row_dimensions[r_num].height = config.DATA_ROW_HEIGHT

            for c_num, val in enumerate(r_data, 1):
                cell = ws.cell(row=r_num, column=c_num, value=val)

                # Fonts & Alignments
                if c_num == 1:
                    cell.alignment = BODY_CENTER_ALIGNMENT
                    cell.font = BODY_FONT
                elif c_num in (2, 3, 4):
                    cell.alignment = BODY_LEFT_ALIGNMENT
                    cell.font = BODY_ITALIC_FONT if (r_offset == 4 and c_num == 4) else BODY_FONT
                elif c_num in (5, 6, 7):
                    cell.alignment = BODY_CENTER_ALIGNMENT
                    cell.font = BODY_FONT

                cell.border = THIN_BORDER

        # Merge Sl No (Column A) across 5 rows
        ws.merge_cells(start_row=start_row, start_column=1, end_row=start_row + 4, end_column=1)

        # Merge Test Name (Column B) across 5 rows
        ws.merge_cells(start_row=start_row, start_column=2, end_row=start_row + 4, end_column=2)

        # Preserve thin borders across all cells in the 5-row block
        apply_range_borders(ws, start_col=1, start_row=start_row, end_col=7, end_row=start_row + 4, border=THIN_BORDER)

        current_row += 5

    # Apply Column Widths
    for col_letter, width in config.COLUMN_WIDTHS.items():
        ws.column_dimensions[col_letter].width = width

    # Page Setup & Print Formatting
    ws.page_setup.orientation = ws.ORIENTATION_LANDSCAPE
    ws.page_setup.paperSize = ws.PAPERSIZE_A4
    ws.page_setup.fitToWidth = 1
    ws.page_setup.fitToHeight = 0
    ws.sheet_properties.pageSetUpPr.fitToPage = True

    # Freeze panes on Header row
    ws.freeze_panes = "A2"

    # Print titles (repeat row 1 on every printed page)
    ws.print_title_rows = "1:1"

    # Show Grid lines
    if ws.views.sheetView:
        ws.views.sheetView[0].showGridLines = True

    logger.info(f"Workbook construction completed successfully. Total rows written: {current_row - 1}")
    return wb


def save_workbook(wb: openpyxl.Workbook, output_path: str) -> str:
    """Saves openpyxl Workbook to specified filepath."""
    wb.save(output_path)
    logger.info(f"Workbook saved to disk: '{output_path}'")
    return output_path
