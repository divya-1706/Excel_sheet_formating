"""
Creates realistic sample OTIS Test Sequence input Excel file.
Saves to sample/sample_test_sequence.xlsx.
"""

import os
import sys

# Ensure parent directory is in sys.path
BASE_DIR = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))
if BASE_DIR not in sys.path:
    sys.path.insert(0, BASE_DIR)

import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
import config


def create_sample_workbook():
    os.makedirs(config.SAMPLE_DIR, exist_ok=True)
    sample_file_path = os.path.join(config.SAMPLE_DIR, "sample_test_sequence.xlsx")

    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = config.DEFAULT_SHEET_NAME

    headers = [
        "Sl.No", "Test Name", "Point", "Description",
        "Voltage", "Current", "Device", "Remarks", "Actual Value"
    ]

    ws.append(headers)

    # Style Header
    header_font = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
    header_fill = PatternFill(start_color="1E293B", end_color="1E293B", fill_type="solid")
    thin_border = Border(
        left=Side(style='thin', color='CBD5E1'),
        right=Side(style='thin', color='CBD5E1'),
        top=Side(style='thin', color='CBD5E1'),
        bottom=Side(style='thin', color='CBD5E1')
    )

    for col_idx, h in enumerate(headers, 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.font = header_font
        cell.fill = header_fill
        cell.alignment = Alignment(horizontal="center", vertical="center")
        cell.border = thin_border

    sample_data = [
        # Report Block 1.1
        ["1.1", "PO SERVICE BRAKE TEST", "P_BAT", "CONNECT BATTERY SUPPLY TO PSU.01", 24.0, 1.5, "PSU.01", "OK", ""],
        ["", "", "P_BAT", "SET Voltage 24V", "", "", "PSU.01", "", ""],
        ["", "", "P_BAT", "SET CURRENT 1.5A", "", "", "PSU.01", "", ""],
        ["", "", "P30V", "CONNECT WAVESHARE TO P30V POINT", "", "", "Waveshare", "", ""],
        ["", "", "P30V", "READ THE VALUE IN DMM FOR P30V", "", "", "DMM", "", ""],

        # Report Block 1.2
        ["1.2", "MAIN CONTACTOR CHECK", "P_BAT", "CONNECT BATTERY SUPPLY TO PSU.01", 18.0, 2.0, "PSU.01", "OK", ""],
        ["", "", "P_BAT", "SET Voltage 18V", "", "", "PSU.01", "", ""],
        ["", "", "P_BAT", "SET CURRENT 2.0A", "", "", "PSU.01", "", ""],
        ["", "", "P18V", "CONNECT WAVESHARE TO P18V NODE", "", "", "Waveshare", "", ""],
        ["", "", "P18V", "READ THE VALUE IN DMM FOR P18V", "", "", "DMM", "", ""],

        # Report Block 2.1
        ["2.1", "SAFETY CHAIN VOLTAGE TEST", "P_BAT", "CONNECT BATTERY SUPPLY TO PSU.02", 48.0, 1.0, "PSU.02", "OK", ""],
        ["", "", "P_BAT", "SET Voltage 48V", "", "", "PSU.02", "", ""],
        ["", "", "P_BAT", "SET CURRENT 1.0A", "", "", "PSU.02", "", ""],
        ["", "", "P48V", "MEASURE WAVESHARE SIGNAL ON P48V", "", "", "Waveshare", "", ""],
        ["", "", "P48V", "READ THE VALUE IN DMM FOR P48V", "", "", "DMM", "", ""],

        # Report Block 2.2
        ["2.2", "DOOR OPERATOR VOLTAGE SENSING", "P_BAT", "CONNECT BATTERY SUPPLY TO PSU.02", 60.0, 3.0, "PSU.02", "OK", ""],
        ["", "", "P_BAT", "SET Voltage 60V", "", "", "PSU.02", "", ""],
        ["", "", "P_BAT", "SET CURRENT 3.0A", "", "", "PSU.02", "", ""],
        ["", "", "P60V", "CONNECT WAVESHARE CHANNEL TO P60V", "", "", "Waveshare", "", ""],
        ["", "", "P60V", "READ THE VALUE IN DMM FOR P60V", "", "", "DMM", "", ""],

        # Report Block 3.1
        ["3.1", "AUXILIARY POWER SUPPLY TEST", "P_BAT", "CONNECT BATTERY SUPPLY TO PSU.01", 30.0, 1.2, "PSU.01", "OK", ""],
        ["", "", "P_BAT", "SET Voltage 30V", "", "", "PSU.01", "", ""],
        ["", "", "P_BAT", "SET CURRENT 1.2A", "", "", "PSU.01", "", ""],
        ["", "", "P30V", "CHECK WAVESHARE READING ON P30V", "", "", "Waveshare", "", ""],
        ["", "", "P30V", "READ THE VALUE IN DMM FOR P30V", "", "", "DMM", "", ""],
    ]

    for row_idx, r in enumerate(sample_data, 2):
        for col_idx, val in enumerate(r, 1):
            cell = ws.cell(row=row_idx, column=col_idx, value=val)
            cell.font = Font(name="Calibri", size=10)
            cell.border = thin_border

    # Set column widths
    widths = [10, 35, 15, 40, 12, 12, 15, 12, 15]
    for i, w in enumerate(widths, 1):
        col_letter = openpyxl.utils.get_column_letter(i)
        ws.column_dimensions[col_letter].width = w

    wb.save(sample_file_path)
    print(f"Sample workbook created successfully at: {sample_file_path}")


if __name__ == "__main__":
    create_sample_workbook()
