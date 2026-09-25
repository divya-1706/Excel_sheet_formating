r"""
Parser module for OTIS Excel Report Formatter.
Reads OTIS Test Sequence Excel files (including 23-column template formats),
detects report blocks from Column B (Sl.No X.Y) or Column A,
and extracts testing parameters (report_no, test_name, battery_point, voltage, current, waveshare_point, dmm_point, device, instructions).
"""

from dataclasses import dataclass, field
import io
import re
from typing import List, Dict, Any, Optional, Union, Tuple
import openpyxl
from openpyxl.worksheet.worksheet import Worksheet

import config
from logger import logger
from utils import clean_str, extract_voltage, extract_current, get_voltage_min_max, extract_number


@dataclass
class ReportBlock:
    """
    Represents a single parsed OTIS report sequence block.
    """
    report_no: str
    test_name: str
    voltage: Optional[float] = None
    voltage_raw: str = ""
    current: Optional[float] = None
    current_raw: str = ""
    waveshare_point: str = ""
    dmm_point: str = ""
    battery_point: str = "P_BAT"
    device: str = config.DEFAULT_PSU_DEVICE
    psu_device: str = config.DEFAULT_PSU_DEVICE
    waveshare_device: str = config.DEFAULT_WAVESHARE_DEVICE
    dmm_device: str = config.DEFAULT_DMM_DEVICE
    instructions: List[str] = field(default_factory=list)
    raw_rows: List[Dict[str, Any]] = field(default_factory=list)
    warnings: List[str] = field(default_factory=list)

    @property
    def measurement_point(self) -> str:
        return self.waveshare_point if self.waveshare_point else "Measurement Point"

    @measurement_point.setter
    def measurement_point(self, value: str):
        self.waveshare_point = value

    @property
    def min_voltage(self) -> Optional[float]:
        min_v, _ = get_voltage_min_max(self.voltage)
        return min_v

    @property
    def max_voltage(self) -> Optional[float]:
        _, max_v = get_voltage_min_max(self.voltage)
        return max_v

    def to_dict(self) -> Dict[str, Any]:
        """Converts dataclass to dictionary for JSON summary rendering."""
        return {
            "report_no": self.report_no,
            "test_name": self.test_name,
            "voltage": self.voltage,
            "voltage_raw": self.voltage_raw,
            "current": self.current,
            "current_raw": self.current_raw,
            "waveshare_point": self.waveshare_point,
            "dmm_point": self.dmm_point,
            "battery_point": self.battery_point,
            "device": self.device,
            "psu_device": self.psu_device,
            "waveshare_device": self.waveshare_device,
            "dmm_device": self.dmm_device,
            "min_voltage": self.min_voltage,
            "max_voltage": self.max_voltage,
            "warnings": self.warnings
        }


def find_test_sequence_sheet(wb: openpyxl.Workbook) -> Worksheet:
    """
    Locates the 'Test Sequence' worksheet in the workbook.
    Case-insensitive fallback matching if exact match isn't found.
    """
    exact_target = config.DEFAULT_SHEET_NAME.lower()
    for sheet_name in wb.sheetnames:
        if sheet_name.strip().lower() == exact_target:
            return wb[sheet_name]

    # Substring search fallback (e.g. "Test Sequence Report", "Sequence")
    for sheet_name in wb.sheetnames:
        if "test" in sheet_name.lower() and "seq" in sheet_name.lower():
            logger.warning(f"Exact sheet 'Test Sequence' not found. Using matched sheet: '{sheet_name}'")
            return wb[sheet_name]

    available = ", ".join([f"'{s}'" for s in wb.sheetnames])
    raise ValueError(
        f"Worksheet '{config.DEFAULT_SHEET_NAME}' not found in workbook. "
        f"Available worksheets: [{available}]"
    )


def detect_headers(ws: Worksheet, max_search_rows: int = 30) -> Tuple[int, Dict[str, int]]:
    """
    Detects header row index and maps column keywords to 1-based column indices.
    Returns (header_row_idx, column_map).
    """
    best_row_idx = -1
    max_matches = 0
    best_row_map: Dict[str, int] = {}

    for row_idx in range(1, max_search_rows + 1):
        row_map: Dict[str, int] = {}
        matches_count = 0
        for col_idx in range(1, ws.max_column + 1):
            cell_val = clean_str(ws.cell(row=row_idx, column=col_idx).value).lower()
            if not cell_val:
                continue

            for col_key, keywords in config.COLUMN_KEYWORDS.items():
                if col_key not in row_map:
                    for kw in keywords:
                        pattern = r"\b" + re.escape(kw) + r"\b"
                        if re.search(pattern, cell_val, re.IGNORECASE) or kw == cell_val:
                            row_map[col_key] = col_idx
                            matches_count += 1
                            break

        if matches_count > max_matches:
            max_matches = matches_count
            best_row_idx = row_idx
            best_row_map = row_map

    if best_row_idx == -1 or max_matches < 2:
        logger.warning("Header detection heuristics yielded weak matches. Falling back to standard column indices.")
        best_row_idx = 1
        best_row_map = {
            "sl_no": 2,
            "test_name": 3,
            "point": 7,
            "description": 4,
            "voltage": 10,
            "current": 11,
            "device": 12,
            "remarks": 13,
            "actual_value": 14
        }
    else:
        if "sl_no" not in best_row_map:
            best_row_map["sl_no"] = 2
        if "test_name" not in best_row_map:
            best_row_map["test_name"] = 3

    logger.info(f"Header row detected at row {best_row_idx} with mapping: {best_row_map}")
    return best_row_idx, best_row_map


def _is_report_start_col_b(val_a: Any, val_b: Any) -> Tuple[bool, str]:
    """
    Checks if Column B (or Column A) marks the start of a report block matching pattern Sl.No X.Y or X.Y.
    Returns (is_start, report_number).
    """
    for val in (val_b, val_a):
        if val is None:
            continue
        s = clean_str(val)
        if not s:
            continue

        # Ignore header strings
        if s.lower() in ("sl.no", "sl no", "sl_no", "s.no", "sl. no.", "sl.no.", "sl. no", "report no", "report_no"):
            continue

        # Match Sl.No 1.1, Sl.No 1.2, 1.1, 2.1, 10.3, etc.
        match = re.search(r"(?:sl\.?\s*no\.?\s*|s\.?\s*no\.?\s*)?(\d+\.\d+)", s, re.IGNORECASE)
        if match:
            return True, match.group(1)

    return False, ""


# Backward compatibility aliases
_is_report_start = lambda col_a_val: _is_report_start_col_b(col_a_val, None)[0]
_is_new_report_row = _is_report_start


def _extract_report_number(val: Any) -> str:
    """Extracts numeric report number like '1.1' from input value."""
    if not val:
        return ""
    s = clean_str(val)
    match = re.search(r"\b(\d+\.\d+)\b", s)
    if match:
        return match.group(1)
    return s


def parse_report_block(raw_rows: List[Dict[str, Any]]) -> Optional[ReportBlock]:
    """
    Parses a group of raw rows belonging to a single report block.
    Extracts report_no, test_name, voltage, current, waveshare_point, dmm_point, device, instructions.
    """
    if not raw_rows:
        return None

    first_row = raw_rows[0]
    raw_b = first_row.get("col_b", first_row.get("sl_no", ""))
    raw_a = first_row.get("col_a", "")
    is_start, report_no = _is_report_start_col_b(raw_a, raw_b)

    if not report_no:
        report_no = _extract_report_number(raw_b) or _extract_report_number(raw_a)
    if not report_no:
        return None

    # Test Name from Column C or first non-empty test name in block
    test_name = ""
    for r in raw_rows:
        tn = r.get("col_c", r.get("test_name", ""))
        if tn:
            test_name = tn
            break

    block = ReportBlock(
        report_no=report_no,
        test_name=test_name,
        raw_rows=raw_rows,
        instructions=[]
    )

    detected_voltages: List[float] = []
    detected_currents: List[float] = []
    waveshare_pts: List[str] = []
    dmm_pts: List[str] = []
    battery_pts: List[str] = []
    devices: List[str] = []

    for r in raw_rows:
        col_d = r.get("col_d", r.get("description", ""))
        col_e = r.get("col_e", "")
        col_g = r.get("col_g", r.get("point", ""))
        col_h = r.get("col_h", "")
        col_j = r.get("col_j", r.get("voltage", ""))
        col_k = r.get("col_k", r.get("current", ""))
        col_l = r.get("col_l", r.get("device", ""))

        all_text = f"{col_d} {col_e} {col_g} {col_h} {col_l}".strip()

        if col_d:
            block.instructions.append(col_d)
        if col_e and col_e != col_d:
            block.instructions.append(col_e)
        if col_l:
            devices.append(col_l)

        # 1. Voltage from Column J or "SET Voltage" text
        v_j = extract_voltage(col_j)
        if v_j is not None:
            detected_voltages.append(v_j)
        else:
            set_v_match = re.search(r"SET\s+Voltage\s*[:=]?\s*([0-9]+(?:\.[0-9]+)?)", all_text, re.IGNORECASE)
            if set_v_match:
                try:
                    detected_voltages.append(float(set_v_match.group(1)))
                except ValueError:
                    pass

        # 2. Current from Column K or "SET CURRENT" text
        i_k = extract_current(col_k)
        if i_k is not None:
            detected_currents.append(i_k)
        else:
            set_i_match = re.search(r"SET\s+CURRENT\s*[:=]?\s*([0-9]+(?:\.[0-9]+)?)", all_text, re.IGNORECASE)
            if set_i_match:
                try:
                    detected_currents.append(float(set_i_match.group(1)))
                except ValueError:
                    pass

        # 3. Waveshare point from Column D starting with Q (e.g. Q25.3, Q10.5, Q14.2)
        q_match = re.search(r"\b(Q\d+(?:\.\d+)?)\b", f"{col_d} {col_e}", re.IGNORECASE)
        if q_match:
            waveshare_pts.append(q_match.group(1).upper())

        # Battery point from Column G / H
        for pt_val in (col_g, col_h):
            if pt_val:
                pt_upper = pt_val.upper()
                if "BAT" in pt_upper or "P_BAT" in pt_upper:
                    battery_pts.append(pt_val)
                elif not waveshare_pts and not pt_upper.startswith("SET"):
                    waveshare_pts.append(pt_val)

        # 4. DMM point from rows where Column E / text contains "READ THE VALUE IN DMM" or "read voltage"
        if re.search(r"READ\s+THE\s+VALUE\s+IN\s+DMM|read\s+voltage", all_text, re.IGNORECASE):
            pt_dmm_match = re.search(r"\b(P_?[A-Z0-9]+|Q\d+(?:\.\d+)?|TP\d+)\b", all_text, re.IGNORECASE)
            if pt_dmm_match:
                dmm_pts.append(pt_dmm_match.group(1).upper())
            elif col_g:
                dmm_pts.append(col_g)
            elif col_h:
                dmm_pts.append(col_h)

    # Resolve devices
    psu_dev = config.DEFAULT_PSU_DEVICE
    for d in devices:
        if "psu" in d.lower():
            psu_dev = d
            break
        elif d and "dmm" not in d.lower() and "waveshare" not in d.lower():
            psu_dev = d

    block.psu_device = psu_dev
    block.device = psu_dev

    for d in devices:
        if "waveshare" in d.lower():
            block.waveshare_device = d
        elif "dmm" in d.lower():
            block.dmm_device = d

    # Resolve Voltage & Current
    if detected_voltages:
        block.voltage = detected_voltages[0]
        block.voltage_raw = f"{block.voltage}V"
    else:
        block.warnings.append(f"Report {report_no}: Voltage value not detected.")

    if detected_currents:
        block.current = detected_currents[0]
        block.current_raw = f"{block.current}A"

    # Resolve Points
    block.waveshare_point = waveshare_pts[0] if waveshare_pts else "Measurement Point"
    block.dmm_point = dmm_pts[0] if dmm_pts else (block.waveshare_point if block.waveshare_point != "Measurement Point" else "P30V")
    block.battery_point = battery_pts[0] if battery_pts else "P_BAT"

    logger.info(f"Found report: {block.report_no}")
    print(f"Found report: {block.report_no}")

    return block


def parse_test_sequence_workbook(file_input: Union[str, io.BytesIO, bytes]) -> Tuple[List[ReportBlock], List[str], Dict[str, Any]]:
    """
    Main entry point for parsing OTIS Test Sequence workbook.
    Detects report blocks from Column B pattern Sl.No X.Y.
    """
    logger.info("Opening workbook for OTIS parsing...")

    if isinstance(file_input, bytes):
        file_input = io.BytesIO(file_input)

    wb = openpyxl.load_workbook(file_input, data_only=True)
    ws = find_test_sequence_sheet(wb)

    header_row_idx, col_map = detect_headers(ws)

    report_blocks: List[ReportBlock] = []
    global_warnings: List[str] = []
    seen_report_nos: set = set()

    current_raw_block: List[Dict[str, Any]] = []

    for r_idx in range(header_row_idx + 1, ws.max_row + 1):
        col_a_val = ws.cell(row=r_idx, column=1).value
        col_b_val = ws.cell(row=r_idx, column=2).value
        col_c_val = ws.cell(row=r_idx, column=3).value
        col_d_val = ws.cell(row=r_idx, column=4).value
        col_e_val = ws.cell(row=r_idx, column=5).value
        col_f_val = ws.cell(row=r_idx, column=6).value
        col_g_val = ws.cell(row=r_idx, column=7).value
        col_h_val = ws.cell(row=r_idx, column=8).value
        col_i_val = ws.cell(row=r_idx, column=9).value
        col_j_val = ws.cell(row=r_idx, column=10).value
        col_k_val = ws.cell(row=r_idx, column=11).value
        col_l_val = ws.cell(row=r_idx, column=12).value

        sl_no_str = clean_str(col_b_val) or clean_str(col_a_val)
        test_name_str = clean_str(col_c_val)
        desc_d = clean_str(col_d_val)
        desc_e = clean_str(col_e_val)
        point_g = clean_str(col_g_val)
        point_h = clean_str(col_h_val)
        voltage_val = col_j_val
        current_val = col_k_val
        device_str = clean_str(col_l_val)

        # Ignore totally blank rows
        if not any([sl_no_str, test_name_str, desc_d, desc_e, point_g, point_h, str(voltage_val or ''), str(current_val or ''), device_str]):
            continue

        row_dict = {
            "row_idx": r_idx,
            "col_a": clean_str(col_a_val),
            "col_b": clean_str(col_b_val),
            "col_c": clean_str(col_c_val),
            "col_d": desc_d,
            "col_e": desc_e,
            "col_g": point_g,
            "col_h": point_h,
            "col_j": voltage_val,
            "col_k": current_val,
            "col_l": device_str,
            "sl_no": sl_no_str,
            "test_name": test_name_str,
            "point": point_g or point_h,
            "description": f"{desc_d} {desc_e}".strip(),
            "voltage": voltage_val,
            "current": current_val,
            "device": device_str
        }

        is_start, _ = _is_report_start_col_b(col_a_val, col_b_val)
        if is_start:
            if current_raw_block:
                parsed_block = parse_report_block(current_raw_block)
                if parsed_block:
                    if parsed_block.report_no in seen_report_nos:
                        w_msg = f"Duplicate Report Number '{parsed_block.report_no}' detected."
                        global_warnings.append(w_msg)
                        parsed_block.warnings.append(w_msg)
                    else:
                        seen_report_nos.add(parsed_block.report_no)
                    report_blocks.append(parsed_block)

            current_raw_block = [row_dict]
        else:
            if current_raw_block:
                current_raw_block.append(row_dict)
            else:
                logger.warning(f"Row {r_idx} skipped (no preceding Sl.No detected): '{sl_no_str}'")

    # Process final block
    if current_raw_block:
        parsed_block = parse_report_block(current_raw_block)
        if parsed_block:
            if parsed_block.report_no in seen_report_nos:
                w_msg = f"Duplicate Report Number '{parsed_block.report_no}' detected."
                global_warnings.append(w_msg)
                parsed_block.warnings.append(w_msg)
            else:
                seen_report_nos.add(parsed_block.report_no)
            report_blocks.append(parsed_block)

    logger.info(f"Detected {len(report_blocks)} report blocks.")
    print(f"Detected {len(report_blocks)} report blocks.")

    summary = {
        "total_reports": len(report_blocks),
        "total_rows_parsed": ws.max_row - header_row_idx,
        "warnings_count": len(global_warnings),
        "header_row": header_row_idx
    }

    return report_blocks, global_warnings, summary
